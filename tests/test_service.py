"""Сервис: две выгрузки → приоритет, блокировка, сбой Google, Яндекс Диск, демо-данные."""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

import pytest
from conftest import EXPIRY_COLUMNS, STOCK_COLUMNS, write_xlsx
from openpyxl import load_workbook

from pharmacy_bot.demo import build_demo
from pharmacy_bot.reports.models import ReportKind
from pharmacy_bot.service import BusyError, ReportService, ServiceError
from pharmacy_bot.yadisk import DiskFile


@pytest.fixture
def files(tmp_path: Path) -> tuple[Path, Path]:
    stock = write_xlsx(
        tmp_path / "остатки.xlsx",
        [
            ("Аптека №1", "Омепразол", "A1", 10, 4000, "01.05.2026"),
            ("Аптека №2", "Лоратадин", "C1", 5, 500, "01.09.2026"),
        ],
        STOCK_COLUMNS,
    )
    expiry = write_xlsx(
        tmp_path / "сроки.xlsx",
        [
            ("Аптека №1", "Омепразол", "A1", "20.09.2026", 10, 4000),
            ("Аптека №2", "Лоратадин", "C1", "01.12.2027", 5, 500),
        ],
        EXPIRY_COLUMNS,
    )
    return stock, expiry


async def test_second_file_completes_priority(service: ReportService, files: tuple[Path, Path]) -> None:
    stock, expiry = files

    first = await service.process_file(stock, "остатки.xlsx")
    second = await service.process_file(expiry, "сроки.xlsx")

    assert first.kind is ReportKind.STOCK
    assert "1 позиция на 4 000 ₽ в 1 аптеке" in first.text
    assert "Приоритет посчитается" in first.text
    assert second.kind is ReportKind.EXPIRY
    assert "Приоритет: 1 серия без продаж" in second.text
    workbook = load_workbook(second.report_path)
    assert workbook.sheetnames == ["Неликвид", "Сроки годности", "Приоритет"]
    assert workbook["Приоритет"]["C2"].value == "A1"


async def test_busy_service_refuses_second_file(service: ReportService, files: tuple[Path, Path]) -> None:
    async with service._lock:
        with pytest.raises(BusyError):
            await service.process_file(files[0], "остатки.xlsx")


class BrokenExporter:
    def export(self, *args: Any) -> str:
        raise RuntimeError("квота Google Sheets исчерпана")


async def test_sheet_failure_keeps_the_report(service: ReportService, files: tuple[Path, Path]) -> None:
    service.exporter = BrokenExporter()  # type: ignore[assignment]

    outcome = await service.process_file(files[0], "остатки.xlsx")

    assert outcome.report_path.exists()
    assert outcome.sheet_error == "Google-таблица не обновлена: квота Google Sheets исчерпана"


async def test_disk_is_required_for_latest_file(service: ReportService) -> None:
    with pytest.raises(ServiceError, match="YANDEX_OAUTH_TOKEN"):
        await service.process_latest(ReportKind.STOCK)


class FakeDisk:
    def __init__(self, source: Path | None) -> None:
        self.source = source

    async def latest_file(self, folder: str, suffixes: tuple[str, ...]) -> DiskFile | None:
        if self.source is None:
            return None
        return DiskFile(self.source.name, f"{folder}{self.source.name}", "", self.source.stat().st_size)

    async def download(self, file: DiskFile, target: Path) -> Path:
        assert self.source is not None
        shutil.copy(self.source, target)
        return target


async def test_latest_file_from_disk_is_processed_and_removed(
    service: ReportService, files: tuple[Path, Path]
) -> None:
    service.disk = FakeDisk(files[1])  # type: ignore[assignment]

    outcome = await service.process_latest(ReportKind.EXPIRY)

    assert outcome.source == "Я.Диск: сроки.xlsx"
    assert not [path for path in service.workdir.iterdir() if path.is_file()]


async def test_empty_disk_folder_is_explained(service: ReportService) -> None:
    service.disk = FakeDisk(None)  # type: ignore[assignment]

    with pytest.raises(ServiceError, match="нет файлов Excel"):
        await service.process_latest(ReportKind.STOCK)


async def test_demo_exports_end_to_end(service: ReportService, tmp_path: Path) -> None:
    stock, expiry = build_demo(tmp_path / "examples")

    await service.process_file(stock, stock.name)
    outcome = await service.process_file(expiry, expiry.name)

    assert "Просрочено" in outcome.text
    workbook = load_workbook(outcome.report_path)
    assert all(workbook[name].max_row > 1 for name in workbook.sheetnames)
    assert "Сейчас свободен" in service.status()
