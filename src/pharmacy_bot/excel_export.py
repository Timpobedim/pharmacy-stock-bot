"""Отчёт в Excel, который бот присылает в ответ: три листа, просроченное подсвечено."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pharmacy_bot.reports.models import ExpiryRow, IlliquidRow, PriorityRow

ILLIQUID_HEADERS = ["Аптека", "Товар", "Серии", "Остаток", "Сумма, ₽", "Последняя продажа", "Дней без продаж"]
EXPIRY_HEADERS = ["Аптека", "Товар", "Серия", "Годен до", "Осталось", "Остаток", "Сумма, ₽"]
PRIORITY_HEADERS = ["Аптека", "Товар", "Серия", "Годен до", "Осталось", "Сумма, ₽", "Дней без продаж"]

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
EXPIRED_FILL = PatternFill("solid", fgColor="F8CBAD")
DATE_FORMAT = "DD.MM.YYYY"
MONEY_FORMAT = "#,##0.00"


def _sheet(sheet: Worksheet, headers: list[str], rows: Sequence[Sequence[Any]], widths: list[int]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append(list(row))
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _format_columns(sheet: Worksheet, dates: tuple[int, ...], money: tuple[int, ...]) -> None:
    for row in sheet.iter_rows(min_row=2):
        for index in dates:
            row[index - 1].number_format = DATE_FORMAT
        for index in money:
            row[index - 1].number_format = MONEY_FORMAT


def write_report(
    path: Path, illiquid: list[IlliquidRow], expiry: list[ExpiryRow], priority: list[PriorityRow]
) -> Path:
    workbook = Workbook()
    del workbook[workbook.sheetnames[0]]

    stale = workbook.create_sheet("Неликвид")
    _sheet(
        stale,
        ILLIQUID_HEADERS,
        [
            (
                r.pharmacy,
                r.product,
                "\n".join(r.series),
                r.quantity,
                r.amount,
                r.last_sale,
                r.days_without_sales,
            )
            for r in illiquid
        ],
        [26, 44, 16, 10, 14, 18, 16],
    )
    _format_columns(stale, dates=(6,), money=(5,))

    soon = workbook.create_sheet("Сроки годности")
    _sheet(
        soon,
        EXPIRY_HEADERS,
        [(r.pharmacy, r.product, r.series, r.expires, r.status, r.quantity, r.amount) for r in expiry],
        [26, 44, 14, 12, 12, 10, 14],
    )
    _format_columns(soon, dates=(4,), money=(7,))
    for row, item in zip(soon.iter_rows(min_row=2), expiry, strict=True):
        if item.expired:
            for cell in row:
                cell.fill = EXPIRED_FILL

    first = workbook.create_sheet("Приоритет")
    _sheet(
        first,
        PRIORITY_HEADERS,
        [
            (r.pharmacy, r.product, r.series, r.expires, f"{r.days_left} дн.", r.amount, r.days_without_sales)
            for r in priority
        ],
        [26, 44, 14, 12, 12, 14, 16],
    )
    _format_columns(first, dates=(4,), money=(6,))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
